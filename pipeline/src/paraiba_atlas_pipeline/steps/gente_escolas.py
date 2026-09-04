"""IDEB per município from INEP, anos iniciais and anos finais of ensino fundamental.

IDEB (Índice de Desenvolvimento da Educação Básica) runs 0 to 10 and combines a proficiency
test with a pass-rate indicator. It is published every two years; 2023 is the latest round.

The rede "Pública" row is used, which merges the municipal and state networks, because a
município's schooling is split between the two and neither alone describes the place.

No school point layer is emitted: INEP's Censo Escolar microdata carries no coordinates
(checked 2021 and 2024, 370 and 426 columns, no latitude or longitude field), and the
catálogo de escolas and geolocalização download paths return 404.
"""
import base64
import json
import textwrap
import time
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import certifi
import openpyxl
import requests

from ..metrics import emit_metric
from ..paths import RAW_DIR
from ..sidra import sidra_values

YEAR = 2023
LICENSE = "INEP open data (Ministério da Educação)"
BASE = "https://download.inep.gov.br/ideb/resultados/divulgacao_{stage}_municipios_{year}.zip"
UF = "PB"
REDE = "Pública"
SCORE_COLUMN = f"VL_OBSERVADO_{YEAR}"
CODE_ROW = 10  # row 10 holds the machine-readable column codes; the rows above are a printed banner

STAGES = {
    "anos_iniciais": ("Anos iniciais", "1º ao 5º ano do ensino fundamental"),
    "anos_finais": ("Anos finais", "6º ao 9º ano do ensino fundamental"),
}


INTERMEDIATE_CRT = "http://secure.globalsign.com/cacert/rnpicpedugr46ovtlsca2025.crt"


def trust_bundle() -> Path:
    """download.inep.gov.br serves only its leaf certificate, omitting the RNP intermediate.

    The intermediate is published at the leaf's Authority Information Access URL and chains up
    to a GlobalSign root that certifi already trusts, so completing the chain keeps verification
    on rather than turning it off.
    """
    bundle = RAW_DIR / "inep_chain.pem"
    if bundle.exists():
        return bundle
    RAW_DIR.mkdir(parents=True, exist_ok=True)
    der = requests.get(INTERMEDIATE_CRT, timeout=60).content
    pem = "-----BEGIN CERTIFICATE-----\n" + "\n".join(textwrap.wrap(base64.b64encode(der).decode(), 64)) + "\n-----END CERTIFICATE-----\n"
    bundle.write_text(Path(certifi.where()).read_text() + "\n" + pem)
    return bundle


def download_resumable(name: str, url: str, *, attempts: int = 8) -> Path:
    """INEP truncates long downloads mid-stream, so resume until the file matches Content-Length."""
    target = RAW_DIR / name
    sidecar = target.with_suffix(target.suffix + ".provenance.json")
    if target.exists() and sidecar.exists():
        return target

    RAW_DIR.mkdir(parents=True, exist_ok=True)
    session = requests.Session()
    session.headers["User-Agent"] = "paraiba-atlas-pipeline"
    session.verify = str(trust_bundle())

    expected = None
    for _ in range(attempts):
        try:
            expected = int(session.head(url, timeout=60, allow_redirects=True).headers["Content-Length"])
            break
        except requests.RequestException:
            time.sleep(3)
    if expected is None:
        raise RuntimeError(f"{name}: could not reach {url}")

    for _ in range(attempts):
        have = target.stat().st_size if target.exists() else 0
        if have >= expected:
            break
        try:
            with session.get(url, headers={"Range": f"bytes={have}-"}, stream=True, timeout=300) as response:
                response.raise_for_status()
                with open(target, "ab") as out:
                    for chunk in response.iter_content(1 << 20):
                        out.write(chunk)
        except requests.RequestException:
            time.sleep(3)
    if not target.exists() or target.stat().st_size != expected:
        raise RuntimeError(f"{name}: got {target.stat().st_size if target.exists() else 0} of {expected} bytes")

    sidecar.write_text(json.dumps({
        "url": url,
        "fetched_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "bytes": expected,
        "license": LICENSE,
    }, ensure_ascii=False, indent=2))
    return target


def read_scores(zip_path: Path) -> dict[str, float]:
    """Return the IDEB score per 7-digit IBGE code for the public network in Paraíba."""
    with ZipFile(zip_path) as archive:
        member = next(n for n in archive.namelist() if n.endswith(".xlsx") and not Path(n).name.startswith("~$"))
        # read_only keeps the handle open lazily, so hand openpyxl the bytes rather than the archive member
        workbook = openpyxl.load_workbook(BytesIO(archive.read(member)), read_only=True, data_only=True)

    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(min_row=CODE_ROW, values_only=True)
    codes = [str(code) if code is not None else "" for code in next(rows)]
    column = {name: index for index, name in enumerate(codes)}
    uf, cod, rede, score = column["SG_UF"], column["CO_MUNICIPIO"], column["REDE"], column[SCORE_COLUMN]

    scores: dict[str, float] = {}
    for row in rows:
        if row[uf] != UF or row[rede] != REDE:
            continue
        raw = row[score]
        if raw in (None, "-", ""):
            continue
        scores[str(row[cod])] = float(str(raw).replace(",", "."))
    workbook.close()
    return scores


def run() -> None:
    population, _ = sidra_values("sidra_4709_populacao_2022_pb.json", 4709, 93, "2022")

    for stage, (label, detail) in STAGES.items():
        url = BASE.format(stage=stage, year=YEAR)
        scores = read_scores(download_resumable(f"inep_ideb_{stage}_municipios_{YEAR}.zip", url))
        emit_metric(
            layer_id=f"gente.ideb_{stage}", path=f"gente/ideb_{stage}.json",
            label=f"IDEB, {label.lower()}", unit="nota de 0 a 10", year=YEAR,
            source=f"INEP, IDEB {YEAR}, rede pública, {detail}", source_url=url,
            values=scores, meso_method="pop_weighted_mean", weights=population, higher_is="better",
        )
        ordered = sorted(scores.items(), key=lambda kv: kv[1])
        print(f"ideb {stage}: {len(scores)} de 223 municípios, "
              f"menor {ordered[0][1]:.1f}, maior {ordered[-1][1]:.1f}")

    print("escolas: INEP não publica latitude/longitude no Censo Escolar (2021 e 2024 conferidos); camada de pontos não emitida")
